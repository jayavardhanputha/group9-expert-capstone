from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

OUT = Path(__file__).with_name("Azure_DevSecOps_Inventory_Architecture.pptx")
prs = Presentation()
prs.slide_width = Inches(13.333)
prs.slide_height = Inches(7.5)

NAVY = RGBColor(11, 18, 38)
PANEL = RGBColor(20, 32, 59)
PANEL2 = RGBColor(27, 43, 72)
WHITE = RGBColor(245, 248, 255)
MUTED = RGBColor(166, 183, 211)
CYAN = RGBColor(57, 211, 222)
BLUE = RGBColor(89, 145, 255)
GREEN = RGBColor(71, 210, 157)
AMBER = RGBColor(255, 190, 92)
RED = RGBColor(255, 112, 124)
FONT = "Aptos"


def textbox(slide, x, y, w, h, text, size=16, color=WHITE, bold=False,
            align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.MIDDLE, margin=0.05):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(margin)
    tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    for idx, line in enumerate(text.split("\n")):
        p = tf.paragraphs[0] if idx == 0 else tf.add_paragraph()
        p.text = line
        p.alignment = align
        p.font.name = FONT
        p.font.size = Pt(size)
        p.font.bold = bold
        p.font.color.rgb = color
    return shape


def rect(slide, x, y, w, h, fill=PANEL, line=None, radius=True):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill
    shape.line.fill.background() if line is None else None
    if line is not None:
        shape.line.color.rgb = line
        shape.line.width = Pt(1.2)
    if radius:
        shape.adjustments[0] = 0.12
    return shape


def card(slide, x, y, w, h, title, body="", accent=CYAN, title_size=17, body_size=12):
    rect(slide, x, y, w, h, PANEL, line=PANEL2)
    rect(slide, x, y, 0.07, h, accent, radius=False)
    textbox(slide, x + 0.22, y + 0.12, w - 0.38, 0.38, title, title_size, WHITE, True)
    if body:
        textbox(slide, x + 0.22, y + 0.55, w - 0.4, h - 0.66, body, body_size, MUTED, valign=MSO_ANCHOR.TOP)


def connector(slide, x1, y1, x2, y2, color=CYAN, width=2.0, arrow=True):
    c = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = color
    c.line.width = Pt(width)
    return c


def base(title, kicker="AZURE DEVSECOPS · INVENTORY PLATFORM"):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = NAVY
    # Accent stripe and subtle page footer.
    rect(slide, 0, 0, 0.12, 7.5, CYAN, radius=False)
    textbox(slide, 0.55, 0.28, 11.8, 0.25, kicker, 9, CYAN, True)
    textbox(slide, 0.55, 0.62, 12.1, 0.56, title, 27, WHITE, True)
    textbox(slide, 0.6, 7.12, 11.9, 0.2, "CAPSTONE REFERENCE DESIGN  ·  DEMO SCOPE — NOT PRODUCTION READY", 8, MUTED)
    textbox(slide, 12.35, 7.08, 0.45, 0.25, f"{len(prs.slides):02d}", 9, MUTED, align=PP_ALIGN.RIGHT)
    return slide


def bullet_list(slide, x, y, w, items, size=15, color=WHITE, gap=0.66, dot=CYAN):
    for i, item in enumerate(items):
        yy = y + i * gap
        rect(slide, x, yy + 0.16, 0.09, 0.09, dot, radius=False)
        textbox(slide, x + 0.22, yy, w - 0.22, 0.52, item, size, color, valign=MSO_ANCHOR.TOP)


# 1 — Title
slide = prs.slides.add_slide(prs.slide_layouts[6])
slide.background.fill.solid(); slide.background.fill.fore_color.rgb = NAVY
rect(slide, 0, 0, 0.15, 7.5, CYAN, radius=False)
# Graphic motif
for i, col in enumerate((CYAN, BLUE, GREEN)):
    rect(slide, 9.1 + i * 0.65, 1.05 + i * 0.36, 2.4, 0.13, col, radius=False)
textbox(slide, 0.85, 1.15, 7.8, 0.35, "AZURE DEVSECOPS · CAPSTONE ARCHITECTURE", 11, CYAN, True)
textbox(slide, 0.8, 2.0, 8.0, 1.6, "Inventory Management\nAPI Platform", 34, WHITE, True, valign=MSO_ANCHOR.TOP)
textbox(slide, 0.85, 4.0, 7.5, 0.8, "A secure-by-default delivery path—from pull request to an approval-gated Azure demo.", 19, MUTED, valign=MSO_ANCHOR.TOP)
rect(slide, 0.85, 5.45, 3.1, 0.55, PANEL, line=PANEL2)
textbox(slide, 1.02, 5.56, 2.75, 0.3, "FLASK  ·  GITHUB ACTIONS  ·  TERRAFORM", 10, CYAN, True)
textbox(slide, 0.85, 7.05, 8.0, 0.2, "REFERENCE ARCHITECTURE  |  OCTOBER 2026", 9, MUTED)

# 2 — Purpose and boundary
slide = base("A small system that demonstrates the whole delivery path")
card(slide, 0.7, 1.55, 3.75, 2.05, "Application", "Flask REST API for inventory CRUD\nSQLite persistence for demo use\nContainerized, non-root runtime", CYAN)
card(slide, 4.78, 1.55, 3.75, 2.05, "Engineering workflow", "Pull requests trigger tests, image build, security scans, and Terraform validation\nNo Azure credentials in normal CI", BLUE, body_size=11)
card(slide, 8.86, 1.55, 3.75, 2.05, "Optional Azure demo", "Manual Terraform deployment to Container Apps\nOIDC federation + protected approval\nResources are billable", GREEN, body_size=11)
rect(slide, 0.7, 4.05, 11.9, 1.7, PANEL2)
textbox(slide, 1.0, 4.3, 2.2, 0.4, "DESIGN INTENT", 12, AMBER, True)
textbox(slide, 1.0, 4.78, 10.9, 0.65, "Keep local development and CI useful without an Azure subscription; make cloud provisioning explicit, reviewable, and opt-in.", 19, WHITE, True, valign=MSO_ANCHOR.TOP)
textbox(slide, 0.82, 6.15, 11.5, 0.5, "The repository describes a teaching sample. The supplied specification files began as headings; the implementation is an explicit sample interpretation.", 11, MUTED)

# 3 — Architecture diagram
slide = base("Architecture at a glance")
# Left authoring and CI lane
textbox(slide, 0.72, 1.42, 4.7, 0.28, "CHANGE + CONTINUOUS VALIDATION", 10, CYAN, True)
card(slide, 0.72, 1.85, 1.55, 0.9, "Developer", "Code · tests", BLUE, 13, 10)
card(slide, 2.58, 1.85, 1.55, 0.9, "GitHub PR", "Review gate", BLUE, 13, 10)
card(slide, 4.4, 1.85, 2.05, 1.05, "GitHub Actions", "API tests · build\nBandit · pip-audit\nTerraform · docs", CYAN, 14, 10)
connector(slide, 2.27, 2.3, 2.55, 2.3, BLUE)
connector(slide, 4.13, 2.3, 4.36, 2.3, CYAN)
# deployment control
card(slide, 0.72, 3.55, 1.9, 1.1, "Operator", "Manual dispatch", AMBER, 14, 11)
card(slide, 3.0, 3.55, 1.9, 1.1, "GitHub OIDC", "Short-lived identity", AMBER, 14, 11)
card(slide, 5.25, 3.55, 1.9, 1.1, "Terraform", "Plan → approval → apply", AMBER, 14, 10)
connector(slide, 2.63, 4.1, 2.97, 4.1, AMBER)
connector(slide, 4.92, 4.1, 5.22, 4.1, AMBER)
# Cloud boundary
rect(slide, 7.55, 1.45, 5.0, 4.65, RGBColor(15, 30, 48), line=CYAN)
textbox(slide, 7.85, 1.65, 4.35, 0.3, "AZURE SUBSCRIPTION · OPTIONAL DEMO", 10, CYAN, True)
card(slide, 8.0, 2.2, 4.1, 1.1, "Azure Container Apps", "Public HTTP ingress · single replica · health probes", GREEN, 15, 11)
card(slide, 8.0, 3.75, 1.9, 1.15, "Azure Files", "SQLite demo data", BLUE, 13, 11)
card(slide, 10.2, 3.75, 1.9, 1.15, "Log Analytics", "Container logs", BLUE, 13, 11)
connector(slide, 9.0, 3.33, 8.95, 3.72, GREEN)
connector(slide, 11.0, 3.33, 11.1, 3.72, GREEN)
connector(slide, 7.16, 4.1, 7.72, 4.1, AMBER)
textbox(slide, 7.85, 5.35, 4.25, 0.4, "Terraform remote state: separately bootstrapped private Azure Blob", 10, MUTED)
textbox(slide, 0.8, 5.52, 6.35, 0.45, "PR/CI path does not authenticate to Azure or deploy resources.", 12, MUTED, True)

# 4 — API and data
slide = base("Application and runtime data flow")
card(slide, 0.72, 1.65, 2.4, 1.15, "API client", "HTTP requests", BLUE, 16, 13)
card(slide, 4.15, 1.65, 3.35, 1.15, "Flask Inventory API", "CRUD · payload validation\nJSON errors · parameterized SQL", CYAN, 16, 12)
card(slide, 8.75, 1.65, 3.4, 1.15, "SQLite database", "Local: instance/inventory.db\nAzure demo: mounted /data", GREEN, 16, 12)
connector(slide, 3.17, 2.22, 4.1, 2.22, BLUE)
connector(slide, 7.55, 2.22, 8.7, 2.22, GREEN)
card(slide, 1.0, 3.55, 3.1, 1.2, "Inventory resources", "POST /api/v1/items\nGET list / item · PUT · DELETE", BLUE, 15, 12)
card(slide, 5.08, 3.55, 3.1, 1.2, "Health endpoints", "GET /health/live\nGET /health/ready", CYAN, 15, 12)
card(slide, 9.1, 3.55, 3.1, 1.2, "Container runtime", "Gunicorn · non-root\nDocker health check", GREEN, 15, 12)
rect(slide, 0.95, 5.35, 11.25, 0.85, RGBColor(55, 43, 28), line=AMBER)
textbox(slide, 1.2, 5.48, 10.7, 0.55, "Constraint: SQLite + Azure Files and one replica are for a demo only—not a concurrent, highly available production data tier.", 13, AMBER, True)

# 5 — Delivery pipeline
slide = base("Continuous validation shifts risk left")
steps = [
    ("01", "Pull request", "Review change\nRun workflows", BLUE),
    ("02", "API quality", "pytest\nDocker build", CYAN),
    ("03", "Security", "Bandit · pip-audit\nGitleaks", GREEN),
    ("04", "Infrastructure", "Terraform fmt/init/validate\nMarkdown links", AMBER),
    ("05", "Merge", "Protected branch\nAll checks green", BLUE),
]
for i, (num, title, body, accent) in enumerate(steps):
    x = 0.7 + i * 2.52
    rect(slide, x, 2.0, 2.14, 2.0, PANEL, line=PANEL2)
    textbox(slide, x + 0.18, 2.15, 0.48, 0.34, num, 12, accent, True)
    textbox(slide, x + 0.18, 2.57, 1.78, 0.43, title, 16, WHITE, True)
    textbox(slide, x + 0.18, 3.07, 1.78, 0.68, body, 12, MUTED, valign=MSO_ANCHOR.TOP)
    if i < len(steps) - 1:
        connector(slide, x + 2.17, 3.0, x + 2.47, 3.0, accent, 1.6)
rect(slide, 0.9, 4.7, 11.45, 1.0, PANEL2)
textbox(slide, 1.15, 4.87, 10.9, 0.6, "The default pipeline validates and builds; it does not publish an image or provision Azure resources. Deployment is a separate manual release path.", 15, WHITE, True)

# 6 — Release and infrastructure
slide = base("Azure release is manual, federated, and approval-gated")
release = [
    (0.75, "1 · Prepare", "Publish an immutable\ncontainer image URI", BLUE),
    (3.15, "2 · Authenticate", "GitHub OIDC exchanges\nfor short-lived Azure access", CYAN),
    (5.55, "3 · Plan", "Terraform creates a\nreviewable plan artifact", AMBER),
    (7.95, "4 · Approve", "Protected production\nenvironment reviewers", GREEN),
    (10.35, "5 · Apply", "Provision only after\nexplicit approval", BLUE),
]
for x, title, body, accent in release:
    card(slide, x, 2.0, 2.05, 1.55, title, body, accent, 14, 11)
for i in range(4):
    connector(slide, release[i][0] + 2.07, 2.78, release[i + 1][0] - 0.05, 2.78, release[i][3], 1.6)
textbox(slide, 0.85, 4.15, 5.0, 0.35, "Terraform demo resources", 15, CYAN, True)
bullet_list(slide, 0.85, 4.6, 5.55, ["Resource group + Log Analytics", "Storage account + Azure Files share", "Container Apps environment + API app"], 12, gap=0.45)
textbox(slide, 6.8, 4.15, 5.2, 0.35, "Before running a deployment", 15, AMBER, True)
bullet_list(slide, 6.8, 4.6, 5.5, ["Bootstrap private Blob remote state separately", "Check permissions, regional availability, and cost", "Use real subscription variables and review the plan"], 12, gap=0.45, dot=AMBER)

# 7 — Security posture
slide = base("Security posture: meaningful controls, explicit limits")
textbox(slide, 0.85, 1.55, 5.35, 0.35, "IMPLEMENTED IN THE SAMPLE", 12, GREEN, True)
bullet_list(slide, 0.9, 2.0, 5.4, ["Parameterized SQL and strict request validation", "Non-root container; liveness/readiness probes", "Least-permission CI; OIDC instead of long-lived Azure secret", "Dependency, source, and secret scanning", "Manual deployment with protected approval"], 13, gap=0.65, dot=GREEN)
textbox(slide, 6.85, 1.55, 5.35, 0.35, "REQUIRED BEFORE PRODUCTION", 12, RED, True)
bullet_list(slide, 6.9, 2.0, 5.45, ["API authentication, authorization, and rate limiting", "Private networking or API gateway; TLS policy", "Managed database, migration, backup, and recovery", "Key Vault, image signing/SBOM, runtime alerting", "SLOs, load/security testing, and DR plan"], 13, gap=0.65, dot=RED)
rect(slide, 0.9, 5.85, 11.4, 0.6, RGBColor(56, 29, 44), line=RED)
textbox(slide, 1.1, 5.95, 11.0, 0.34, "Public ingress + unauthenticated API are accepted only for isolated demo data; do not expose sensitive records.", 12, RED, True)

# 8 — Governance and operations
slide = base("Governance makes change and operations reviewable")
roles = [
    ("Developer", "API, tests, compatibility", BLUE),
    ("Platform engineer", "Terraform, workflows, federation", CYAN),
    ("Security reviewer", "Scans, threat model, exceptions", GREEN),
    ("Release approver", "Plan review + protected apply", AMBER),
    ("Service owner", "Availability, data, incidents", BLUE),
]
for i, (title, body, accent) in enumerate(roles):
    x = 0.75 + (i % 3) * 4.1
    y = 1.65 + (i // 3) * 1.35
    card(slide, x, y, 3.65, 1.05, title, body, accent, 15, 11)
card(slide, 0.75, 4.55, 5.55, 1.25, "Change control", "PR review · required checks · architecture decisions recorded", CYAN, 16, 12)
card(slide, 6.85, 4.55, 5.55, 1.25, "Operational loop", "Verify readiness · inspect logs · preserve state · rollback via reviewed Terraform release", GREEN, 16, 12)

# 9 — Evolution roadmap
slide = base("A practical path from demo to production")
roadmap = [
    (0.85, "NOW", "Local + CI", "Run API/tests\nValidate Terraform\nFix scanner findings", BLUE),
    (4.55, "NEXT", "Controlled demo", "Bootstrap state\nConstrain OIDC\nDeploy with approval", CYAN),
    (8.25, "THEN", "Production gates", "Identity + private network\nManaged DB + DR\nMonitoring + SLO", GREEN),
]
for x, tag, title, body, accent in roadmap:
    rect(slide, x, 2.0, 3.15, 2.45, PANEL, line=PANEL2)
    textbox(slide, x + 0.22, 2.2, 2.7, 0.3, tag, 10, accent, True)
    textbox(slide, x + 0.22, 2.65, 2.7, 0.42, title, 18, WHITE, True)
    textbox(slide, x + 0.22, 3.2, 2.7, 0.9, body, 13, MUTED, valign=MSO_ANCHOR.TOP)
connector(slide, 4.04, 3.2, 4.48, 3.2, BLUE, 2.2)
connector(slide, 7.74, 3.2, 8.18, 3.2, CYAN, 2.2)
rect(slide, 0.9, 5.15, 11.4, 0.85, PANEL2)
textbox(slide, 1.15, 5.32, 10.9, 0.48, "Production is a separate architecture decision: resolve identity, data, networking, resilience, and operational ownership before migrating real data.", 14, WHITE, True)

# 10 — Closing summary
slide = base("Key takeaways")
card(slide, 0.8, 1.65, 3.7, 2.0, "01 · Build safely", "Small Flask API, tests, container, and Terraform are maintained together.", CYAN, 17, 13)
card(slide, 4.8, 1.65, 3.7, 2.0, "02 · Validate continuously", "CI runs without cloud credentials and catches quality/security issues before deployment.", BLUE, 17, 13)
card(slide, 8.8, 1.65, 3.7, 2.0, "03 · Deploy deliberately", "Cloud changes require manual dispatch, short-lived identity, plan review, and approval.", GREEN, 17, 13)
textbox(slide, 1.05, 4.45, 11.2, 0.7, "A clear, subscription-optional DevSecOps learning path—with production gaps made visible, not hidden.", 20, WHITE, True, align=PP_ALIGN.CENTER)
textbox(slide, 1.0, 5.65, 11.3, 0.45, "Source: repository reference architecture, platform specification, security strategy, governance, and Terraform configuration.", 10, MUTED, align=PP_ALIGN.CENTER)

prs.save(OUT)
print(f"Wrote {OUT}")
